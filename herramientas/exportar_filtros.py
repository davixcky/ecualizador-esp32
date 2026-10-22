from pathlib import Path
import re
from io import BytesIO
import pymupdf as fitz
from pptx import Presentation
from pptx.util import Inches
from pptx.oxml.xmlchemy import OxmlElement
r=Path(__file__).resolve().parents[1];src=r/'docs/latex/presentacion-filtros.tex';d=r/'docs/presentacion';pdf=fitz.open(r/'docs/latex/presentacion-filtros.pdf');s=src.read_text();notes=re.findall(r'\\note\{(.*?)\}\s*\\end\{frame\}',s,re.S)
notes=[n.replace(r'\_', '_') for n in notes]
assert len(pdf)==len(notes)==10
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
for p,n in zip(pdf,notes):
 sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.shapes.add_picture(BytesIO(p.get_pixmap(matrix=fitz.Matrix(2,2)).tobytes('png')),0,0,width=prs.slide_width,height=prs.slide_height);sl.notes_slide.notes_text_frame.text=n
 for link in p.get_links():
  address=link.get('uri') or link.get('file')
  if not address:continue
  from pptx.enum.shapes import MSO_SHAPE
  box=link['from'];sh=sl.shapes.add_shape(MSO_SHAPE.RECTANGLE,int(box.x0/p.rect.width*prs.slide_width),int(box.y0/p.rect.height*prs.slide_height),int(box.width/p.rect.width*prs.slide_width),int(box.height/p.rect.height*prs.slide_height));sh.fill.background();sh.line.fill.background();sh.click_action.hyperlink.address=address
 tr=OxmlElement('p:transition');tr.set('spd','fast');tr.append(OxmlElement('p:fade'));sl._element.append(tr)
prs.save(d/'presentacion-filtros-hoy.pptx')
for p in [r/'simulacion/velxio'/f'filtros-{hz}Hz.vlx' for hz in (100,1000,8000)]:pdf.embfile_add(p.name,p.read_bytes(),filename=p.name)
pdf.save(d/'presentacion-filtros-hoy.pdf',deflate=True)
print('10 diapositivas de filtros; PDF con 3 proyectos adjuntos y PPTX con notas.')
