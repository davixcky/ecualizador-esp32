#!/usr/bin/env python3
"""Exporta el PDF de Beamer a PNG y PowerPoint con notas; editar docs/latex/*.tex."""
from pathlib import Path
import json, re, shutil
import pymupdf as fitz
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from io import BytesIO
def raster_width(page):
    # El esquema del anexo debe conservar detalle al ampliar en PowerPoint.
    return 6000 if 'Anexo: esquema completo' in page.get_text() else 1920

def copiar_enlaces(page, slide, width, height):
    for link in page.get_links():
        address = link.get('uri') or link.get('file')
        if not address:
            continue
        box = link['from']
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
            int(box.x0/page.rect.width*width), int(box.y0/page.rect.height*height),
            int(box.width/page.rect.width*width), int(box.height/page.rect.height*height))
        shape.fill.background()
        shape.line.fill.background()
        shape.click_action.hyperlink.address = address

def copiar_enlaces_internos(pdf, prs):
    # Resolver después de crear todas las diapositivas, incluido el anexo.
    for i, page in enumerate(pdf):
        for link in page.get_links():
            target = link.get('page', -1)
            if link.get('kind') not in (fitz.LINK_GOTO, fitz.LINK_NAMED) or not 0 <= target < len(prs.slides):
                continue
            box = link['from']
            shape = prs.slides[i].shapes.add_shape(MSO_SHAPE.RECTANGLE,
                int(box.x0/page.rect.width*prs.slide_width), int(box.y0/page.rect.height*prs.slide_height),
                int(box.width/page.rect.width*prs.slide_width), int(box.height/page.rect.height*prs.slide_height))
            shape.fill.background(); shape.line.fill.background()
            shape.click_action.target_slide = prs.slides[target]

root = Path(__file__).resolve().parents[1]
deck = root / 'docs/presentacion'
source = (root / 'docs/latex/presentacion.tex').read_text()
notes = re.findall(r'\\note\{(.*?)\}\s*\\end\{frame\}', source, re.S)
notes = [n.replace(r'\_', '_') for n in notes]
pdf = fitz.open(root / 'docs/latex/presentacion.pdf')
assert len(notes) == len(pdf), (len(notes), len(pdf))
for name, target in [('informe.pdf','docs/informe-fase2.pdf'), ('presentacion.pdf','docs/presentacion/presentacion-fase2.pdf'), ('presentacion-notas.pdf','docs/presentacion/presentacion-notas.pdf')]:
    shutil.copyfile(root / 'docs/latex' / name, root / target)
pngdir = deck / 'png'
pngdir.mkdir(exist_ok=True)
for p in pngdir.glob('*.png'): p.unlink()
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
for i, page in enumerate(pdf):
    png = pngdir / f'{i+1:02d}.png'
    page.get_pixmap(matrix=fitz.Matrix(raster_width(page)/page.rect.width,raster_width(page)/page.rect.width)).save(png)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(png),0,0,width=prs.slide_width,height=prs.slide_height)
    copiar_enlaces(page, slide, prs.slide_width, prs.slide_height)
    slide.notes_slide.notes_text_frame.text = notes[i]
    transition = OxmlElement('p:transition')
    transition.set('spd', 'fast')
    transition.set('advClick', '1')
    transition.append(OxmlElement('p:fade'))
    slide._element.append(transition)
copiar_enlaces_internos(pdf, prs)
prs.save(deck / 'presentacion-fase2.pptx')
(deck / 'notas.json').write_text(json.dumps({str(i+1):n for i,n in enumerate(notes)},ensure_ascii=False,indent=2)+'\n')
print(f'{len(pdf)} diapositivas, PNG y PPTX con notas exportados.')

md = root / 'docs/informe-fase2.md'
md.write_text(re.sub(r'src="(?!\.\./img/)', 'src="../img/', md.read_text()))

# Beamer produces the progressive states; PowerPoint fades between them on click.
animated = fitz.open(root / 'docs/latex/presentacion-animada.pdf')
shutil.copyfile(root / 'docs/latex/presentacion-animada.pdf', deck / 'presentacion-animada.pdf')
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
steps = []
for page in animated:
    match = re.search(rf'(\d+)\s*/\s*{len(notes)}\s*$', page.get_text().strip())
    assert match, page.get_text()[-150:]
    frame = int(match[1])
    steps.append(frame)
    pix = page.get_pixmap(matrix=fitz.Matrix(raster_width(page)/page.rect.width, raster_width(page)/page.rect.width))
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(BytesIO(pix.tobytes('png')),0,0,width=prs.slide_width,height=prs.slide_height)
    copiar_enlaces(page, slide, prs.slide_width, prs.slide_height)
    slide.notes_slide.notes_text_frame.text = notes[frame-1]
    transition = OxmlElement('p:transition')
    transition.set('spd','fast')
    transition.set('advClick','1')
    transition.append(OxmlElement('p:fade'))
    slide._element.append(transition)
copiar_enlaces_internos(animated, prs)
prs.save(deck / 'presentacion-animada.pptx')
(deck / 'animacion.json').write_text(json.dumps({'diapositivas':len(notes), 'pasos':len(steps), 'orden':steps, 'avance':'clic, sin temporizador', 'transicion':'fundido'},ensure_ascii=False,indent=2)+'\n')
print(f'Presentación animada: {len(steps)} pasos / {len(notes)} diapositivas.')

# Mantener la misma familia y acento institucional en la copia Word.
from docx import Document
from docx.shared import RGBColor
word_path = root / 'docs/informe-fase2.docx'
word = Document(word_path)
for style in word.styles:
    if style.type == 1:
        style.font.name = 'Red Hat Display'
        if style.name.startswith('Heading') or style.name == 'Title':
            style.font.color.rgb = RGBColor.from_string('A3161A')
word.save(word_path)

# Los proyectos y planos viajan con el PDF, sin depender de un enlace GitHub todavía no publicado.
attachments = list((root/'simulacion/velxio').glob('integrado-*.vlx')) + [root/'simulacion/velxio/INTEGRADO.md'] + list((root/'simulacion/velxio').glob('filtros-*.vlx')) + [root/'simulacion/velxio/led-animado.vlx',root/'hardware/ecualizador.pdf', root/'hardware/pcb-rev-D2.pdf',root/'hardware/PUENTES.md',root/'hardware/una-cara/LEEME.md',root/'simulacion/velxio/DEMO-HOY.md']
for target in [root/'docs/informe-fase2.pdf',deck/'presentacion-fase2.pdf',deck/'presentacion-animada.pdf']:
    output=fitz.open(target)
    for item in attachments:
        output.embfile_add(item.name,item.read_bytes(),filename=item.name,desc='Proyecto de simulación' if item.suffix=='.vlx' else 'Plano técnico D.2')
    tmp=target.with_suffix('.tmp.pdf')
    output.save(tmp,deflate=True)
    output.close()
    tmp.replace(target)
